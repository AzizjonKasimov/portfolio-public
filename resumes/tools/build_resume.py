"""Build a resume's DOCX and PDF from its Markdown source.

Usage (repo root, project venv):
    .venv\\Scripts\\python.exe resumes\\tools\\build_resume.py <source.md> [<source.md> ...]
    .venv\\Scripts\\python.exe resumes\\tools\\build_resume.py --all

Each source sits next to its outputs and shares their stem, for example
`current/global/ai-rag-agent-engineer/Azizjon_Kasimov_Resume_AI_RAG_Agent_Engineer.md` builds the
`.docx` and `.pdf` beside it. The source format is documented in
`resumes/templates/resume-template.md`; `--all` builds every source under `resumes/current/`.

The DOCX is written with python-docx and exported to PDF by Microsoft Word
(`export_docx_to_pdf.ps1`), so both files always match. The builder keeps the largest body size
between 8.5 and 10 pt that fits the page budget (front matter `pages`, default 1), then runs
`check_resume_pdf.py` and exits non-zero if the PDF breaks a rule. Requires Word on Windows.
"""

import argparse
import re
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

import fitz
from docx import Document
from docx.enum.text import WD_LINE_SPACING, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

TOOLS = Path(__file__).resolve().parent
sys.path.insert(0, str(TOOLS))
from check_resume_pdf import LINK_PATTERNS, check, expected_uri  # noqa: E402

EXPORT_SCRIPT = TOOLS / "export_docx_to_pdf.ps1"
CURRENT = TOOLS.parent / "current"

# A4, because Korean recruiters print and forward A4.
PAGE_WIDTH = 8.27
PAGE_HEIGHT = 11.69
MARGIN_X = 0.6
MARGIN_TOP = 0.48
MARGIN_BOTTOM = 0.42
CONTENT_WIDTH = PAGE_WIDTH - 2 * MARGIN_X

FONTS = {"en": "Arial", "ko": "Malgun Gothic"}
# Each font's single-line height in ems; Word's line-spacing multiple scales it.
SINGLE_LINE = {"en": 1.117, "ko": 1.33}
# Default multiples give about 1.3x leading. Only when the resume still overflows at the smallest
# body size does the builder tighten the spacing, never below LEADING_FLOOR x the font size.
LINE = {"en": 1.15, "ko": 1.0}
LEADING_FLOOR = 1.2
NAVY = RGBColor(0x1B, 0x36, 0x5D)
ACCENT = RGBColor(0x1F, 0x6F, 0x8B)
MUTED = RGBColor(0x55, 0x5F, 0x6B)
BODY = RGBColor(0x1F, 0x24, 0x28)
RULE = "A9C3CF"

# Every size and gap below is designed at a 9 pt body and scales with the chosen body size.
BASE = 9.0
SIZES = [8.5, 9.0, 9.5, 10.0]  # half-point steps: Word stores font sizes in half-points
FILL_TARGET = 0.95  # stop growing once the text reaches this share of the usable page height

LINK_MD = re.compile(r"\[([^\]]+)\]\(([^)\s]+)\)")
BOLD = re.compile(r"\*\*(.+?)\*\*")
AUTO_LINK = re.compile("|".join(f"(?:{pattern.pattern})" for pattern, _ in LINK_PATTERNS))
FONT_ATTRS = ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia")


# ---------------------------------------------------------------------------- source parsing


def parse_source(path):
    """Return the front matter and the resume structure of a Markdown source."""
    raw = path.read_text(encoding="utf-8").replace("\r\n", "\n")
    raw = re.sub(r"<!--.*?-->", "", raw, flags=re.S)
    meta = {}
    front = re.match(r"---\n(.*?)\n---\n", raw, re.S)
    if front:
        for line in front.group(1).splitlines():
            if ":" in line:
                key, value = line.split(":", 1)
                meta[key.strip()] = value.strip()
        raw = raw[front.end():]

    resume = {"name": None, "headline": None, "header": [], "sections": []}
    section = None
    for line in raw.splitlines():
        text = line.strip()
        if not text:
            continue
        if text.startswith("# "):
            resume["name"] = text[2:].strip()
        elif text.startswith("## "):
            section = {"title": text[3:].strip(), "items": []}
            resume["sections"].append(section)
        elif section is None:
            if resume["headline"] is None:
                resume["headline"] = text
            else:
                resume["header"].append(text)
        elif text.startswith("### "):
            section["items"].append(("role", [part.strip() for part in text[4:].split(" | ")]))
        elif text.startswith("- "):
            section["items"].append(("bullet", text[2:].strip()))
        elif re.fullmatch(r"_(.+)_", text):
            section["items"].append(("meta", text[1:-1].strip()))
        elif labeled := re.fullmatch(r"\*\*(.+?):\*\*\s*(.*)", text):
            section["items"].append(("label", labeled.group(1), labeled.group(2)))
        else:
            section["items"].append(("para", text))

    missing = [key for key in ("output", "lang") if key not in meta]
    if missing or not resume["name"] or not resume["headline"]:
        raise SystemExit(f"{path}: needs front matter {missing or ''} and a '# Name' line followed by a headline")
    if meta["lang"] not in FONTS:
        raise SystemExit(f"{path}: lang must be one of {sorted(FONTS)}")
    return meta, resume


def inline_runs(text):
    """Split inline Markdown into (text, url, bold) runs; url is None for plain text."""
    runs = []
    pos = 0
    for match in LINK_MD.finditer(text):
        runs += _bold_runs(text[pos:match.start()])
        runs.append((match.group(1), match.group(2), False))
        pos = match.end()
    return runs + _bold_runs(text[pos:])


def _bold_runs(text):
    runs = []
    pos = 0
    for match in BOLD.finditer(text):
        runs += _link_runs(text[pos:match.start()], False)
        runs += _link_runs(match.group(1), True)
        pos = match.end()
    return runs + _link_runs(text[pos:], False)


def _link_runs(text, bold):
    """Turn visible emails and LinkedIn, GitHub and github.io addresses into links."""
    runs = []
    pos = 0
    for match in AUTO_LINK.finditer(text):
        token = match.group(0).rstrip(".,;)")
        url = expected_uri(token)
        if not url:
            continue
        if match.start() > pos:
            runs.append((text[pos:match.start()], None, bold))
        runs.append((token, url, bold))
        pos = match.start() + len(token)
    if pos < len(text):
        runs.append((text[pos:], None, bold))
    return runs


# ---------------------------------------------------------------------------- DOCX rendering


class Renderer:
    def __init__(self, meta, body_size, line=None):
        self.lang = meta["lang"]
        self.font = FONTS[self.lang]
        self.line = line or LINE[self.lang]
        self.scale = body_size / BASE
        self.body = body_size
        self.doc = Document()
        self._set_defaults(meta)
        self.num_id = self._add_bullet_numbering()

    def pt(self, size):
        return round(size * self.scale, 2)

    # -- document setup

    def _set_defaults(self, meta):
        doc = self.doc
        section = doc.sections[0]
        section.page_width = Inches(PAGE_WIDTH)
        section.page_height = Inches(PAGE_HEIGHT)
        section.left_margin = section.right_margin = Inches(MARGIN_X)
        section.top_margin = Inches(MARGIN_TOP)
        section.bottom_margin = Inches(MARGIN_BOTTOM)

        props = doc.core_properties
        props.title = meta.get("title", "Azizjon Kasimov - Resume")
        props.author = meta.get("author", "Azizjon Kasimov")
        props.keywords = meta.get("keywords", "")

        # Replace the template's theme fonts so Word never falls back to Calibri.
        r_pr = doc.styles.element.find(qn("w:docDefaults")).find(qn("w:rPrDefault")).find(qn("w:rPr"))
        fonts = r_pr.find(qn("w:rFonts"))
        if fonts is None:
            fonts = OxmlElement("w:rFonts")
            r_pr.insert(0, fonts)
        for attr in list(fonts.attrib):
            del fonts.attrib[attr]
        for attr in FONT_ATTRS:
            fonts.set(qn(attr), self.font)

        normal = doc.styles["Normal"]
        normal.font.name = self.font
        normal.font.size = Pt(self.body)
        normal.font.color.rgb = BODY
        normal.paragraph_format.space_before = Pt(0)
        normal.paragraph_format.space_after = Pt(0)
        # Leave w:wordWrap unset: its "0" value lets Word break Latin and Korean words mid-word.
        if self.lang == "ko":
            # Word otherwise widens the gap between Hangul and Latin letters or digits, and the PDF
            # text layer then reads "30 만" where the source says "30만".
            p_pr = normal.element.get_or_add_pPr()
            for tag in ("w:autoSpaceDE", "w:autoSpaceDN"):
                el = OxmlElement(tag)
                el.set(qn("w:val"), "0")
                p_pr.insert_element_before(
                    el, "w:bidi", "w:adjustRightInd", "w:snapToGrid", "w:spacing", "w:ind",
                    "w:contextualSpacing", "w:mirrorIndents", "w:suppressOverlap", "w:jc",
                    "w:textDirection", "w:textAlignment", "w:textboxTightWrap", "w:outlineLvl",
                    "w:divId", "w:cnfStyle", "w:rPr", "w:sectPr", "w:pPrChange",
                )

    def _add_bullet_numbering(self):
        """Register a bullet list drawn in the body font, not the Symbol font."""
        numbering = self.doc.part.numbering_part.element
        abstract_id = 1 + max(
            (int(el.get(qn("w:abstractNumId"))) for el in numbering.findall(qn("w:abstractNum"))), default=0
        )
        num_id = 1 + max((int(el.get(qn("w:numId"))) for el in numbering.findall(qn("w:num"))), default=0)

        abstract = OxmlElement("w:abstractNum")
        abstract.set(qn("w:abstractNumId"), str(abstract_id))
        multi = OxmlElement("w:multiLevelType")
        multi.set(qn("w:val"), "singleLevel")
        abstract.append(multi)
        level = OxmlElement("w:lvl")
        level.set(qn("w:ilvl"), "0")
        for tag, value in (("w:start", "1"), ("w:numFmt", "bullet"), ("w:lvlText", "•"), ("w:lvlJc", "left")):
            el = OxmlElement(tag)
            el.set(qn("w:val"), value)
            level.append(el)
        p_pr = OxmlElement("w:pPr")
        indent = OxmlElement("w:ind")
        indent.set(qn("w:left"), "250")
        indent.set(qn("w:hanging"), "180")
        p_pr.append(indent)
        level.append(p_pr)
        r_pr = OxmlElement("w:rPr")
        fonts = OxmlElement("w:rFonts")
        for attr in FONT_ATTRS:
            fonts.set(qn(attr), self.font)
        r_pr.append(fonts)
        color = OxmlElement("w:color")
        color.set(qn("w:val"), str(ACCENT))
        r_pr.append(color)
        level.append(r_pr)
        abstract.append(level)
        # Word requires every w:abstractNum to precede the first w:num.
        first_num = numbering.find(qn("w:num"))
        if first_num is not None:
            first_num.addprevious(abstract)
        else:
            numbering.append(abstract)
        num = OxmlElement("w:num")
        num.set(qn("w:numId"), str(num_id))
        ref = OxmlElement("w:abstractNumId")
        ref.set(qn("w:val"), str(abstract_id))
        num.append(ref)
        numbering.append(num)
        return num_id

    # -- low-level helpers

    def run(self, paragraph, text, size, color=BODY, bold=False, italic=False, spacing=0.0):
        run = paragraph.add_run(text)
        run.font.name = self.font
        fonts = run._element.get_or_add_rPr().get_or_add_rFonts()
        for attr in FONT_ATTRS:
            fonts.set(qn(attr), self.font)
        run.font.size = Pt(size)
        run.font.color.rgb = color
        run.bold = bold
        run.italic = italic
        if spacing:
            el = OxmlElement("w:spacing")
            el.set(qn("w:val"), str(round(spacing * 20)))
            # Word rejects rPr children out of schema order, so place it before its successors.
            run._element.get_or_add_rPr().insert_element_before(
                el, "w:w", "w:kern", "w:position", "w:sz", "w:szCs", "w:highlight", "w:u", "w:effect",
                "w:bdr", "w:shd", "w:fitText", "w:vertAlign", "w:rtl", "w:cs", "w:em", "w:lang",
                "w:eastAsianLayout", "w:specVanish", "w:oMath",
            )
        return run

    def link(self, paragraph, text, url, size, bold=False):
        rel_id = paragraph.part.relate_to(
            url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True
        )
        hyperlink = OxmlElement("w:hyperlink")
        hyperlink.set(qn("r:id"), rel_id)
        run = OxmlElement("w:r")
        r_pr = OxmlElement("w:rPr")
        fonts = OxmlElement("w:rFonts")
        for attr in FONT_ATTRS:
            fonts.set(qn(attr), self.font)
        r_pr.append(fonts)
        if bold:
            r_pr.append(OxmlElement("w:b"))
        color = OxmlElement("w:color")
        color.set(qn("w:val"), str(ACCENT))
        r_pr.append(color)
        size_el = OxmlElement("w:sz")
        size_el.set(qn("w:val"), str(round(size * 2)))
        r_pr.append(size_el)
        run.append(r_pr)
        text_el = OxmlElement("w:t")
        text_el.text = text
        text_el.set(qn("xml:space"), "preserve")
        run.append(text_el)
        hyperlink.append(run)
        paragraph._p.append(hyperlink)

    def inline(self, paragraph, text, size, color=BODY, italic=False):
        for chunk, url, bold in inline_runs(text):
            if url:
                self.link(paragraph, chunk, url, size, bold)
            else:
                self.run(paragraph, chunk, size, NAVY if bold else color, bold=bold, italic=italic)

    def paragraph(self, before=0.0, after=0.0, line=None, keep_next=False):
        paragraph = self.doc.add_paragraph()
        fmt = paragraph.paragraph_format
        fmt.space_before = Pt(self.pt(before))
        fmt.space_after = Pt(self.pt(after))
        fmt.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
        fmt.line_spacing = line or self.line
        fmt.keep_with_next = keep_next
        return paragraph

    def right_tab(self, paragraph):
        paragraph.paragraph_format.tab_stops.add_tab_stop(Inches(CONTENT_WIDTH), alignment=WD_TAB_ALIGNMENT.RIGHT)

    # -- resume blocks

    def header(self, resume):
        name = self.paragraph(after=1.0, line=1.0, keep_next=True)
        self.run(name, resume["name"].upper(), self.pt(20.0), NAVY, bold=True, spacing=0.6)
        headline = self.paragraph(after=2.4, line=1.0, keep_next=True)
        self.run(headline, resume["headline"], self.pt(10.4), ACCENT, bold=True)
        for i, line in enumerate(resume["header"]):
            last = i == len(resume["header"]) - 1
            self.inline(self.paragraph(after=1.0 if last else 0.8, line=1.0), line, self.pt(8.6), MUTED)

    def section(self, title):
        paragraph = self.doc.add_paragraph()
        # Border first: python-docx then inserts keepNext and spacing around it in schema order.
        p_pr = paragraph._p.get_or_add_pPr()
        borders = OxmlElement("w:pBdr")
        bottom = OxmlElement("w:bottom")
        for key, value in (("w:val", "single"), ("w:sz", "6"), ("w:space", "1"), ("w:color", RULE)):
            bottom.set(qn(key), value)
        borders.append(bottom)
        p_pr.append(borders)
        fmt = paragraph.paragraph_format
        fmt.space_before = Pt(self.pt(5.5))
        fmt.space_after = Pt(self.pt(2.6))
        fmt.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
        fmt.line_spacing = 1.0
        fmt.keep_with_next = True
        self.run(paragraph, title.upper(), self.pt(9.9), NAVY, bold=True, spacing=0.5)

    def role(self, parts, has_meta):
        heading = self.paragraph(before=3.0, after=0 if has_meta else 1.4, line=1.0, keep_next=True)
        self.right_tab(heading)
        # "### Title | Organization | Dates"; the organization may be empty or left out.
        title, dates = parts[0], parts[-1] if len(parts) > 1 else ""
        org = " | ".join(part for part in parts[1:-1] if part)
        self.run(heading, title, self.pt(9.7), NAVY, bold=True)
        if org:
            # Trailing space keeps org and dates apart in the extracted text layer.
            self.run(heading, "  |  " + org + " ", self.pt(9.7), ACCENT, bold=True)
        elif dates:
            self.run(heading, " ", self.pt(9.7))
        if dates:
            self.run(heading, "\t" + dates, self.pt(8.8), MUTED, bold=True)

    def meta(self, text):
        paragraph = self.paragraph(before=0.4, after=1.4, line=1.0, keep_next=True)
        self.inline(paragraph, text, self.pt(8.6), MUTED, italic=self.lang == "en")

    def bullet(self, text):
        paragraph = self.paragraph(after=1.6)
        p_pr = paragraph._p.get_or_add_pPr()
        num_pr = OxmlElement("w:numPr")
        ilvl = OxmlElement("w:ilvl")
        ilvl.set(qn("w:val"), "0")
        num_el = OxmlElement("w:numId")
        num_el.set(qn("w:val"), str(self.num_id))
        num_pr.append(ilvl)
        num_pr.append(num_el)
        p_pr.append(num_pr)
        self.inline(paragraph, text, self.body)

    def label(self, label, value):
        paragraph = self.paragraph(after=1.2)
        self.run(paragraph, label + ": ", self.body, NAVY, bold=True)
        self.inline(paragraph, value, self.body)

    def para(self, text):
        self.inline(self.paragraph(after=0.8), text, self.body)

    def render(self, resume, out_path):
        self.header(resume)
        for section in resume["sections"]:
            self.section(section["title"])
            items = section["items"]
            for i, item in enumerate(items):
                kind = item[0]
                if kind == "role":
                    has_meta = i + 1 < len(items) and items[i + 1][0] == "meta"
                    self.role(item[1], has_meta)
                elif kind == "meta":
                    self.meta(item[1])
                elif kind == "bullet":
                    self.bullet(item[1])
                elif kind == "label":
                    self.label(item[1], item[2])
                else:
                    self.para(item[1])
        self.doc.save(out_path)


# ---------------------------------------------------------------------------- build and fit


def export_pdf(docx_path, pdf_path):
    try:
        result = subprocess.run(
            ["powershell.exe", "-NoProfile", "-ExecutionPolicy", "Bypass", "-File", str(EXPORT_SCRIPT),
             "-InputDocx", str(docx_path), "-OutputPdf", str(pdf_path)],
            capture_output=True, text=True, timeout=180,
        )
    except subprocess.TimeoutExpired:
        raise SystemExit("Word did not finish exporting within 3 minutes. Close any stuck Word "
                         "window or WINWORD.EXE process and run the build again.") from None
    if result.returncode != 0 or not pdf_path.exists():
        raise SystemExit(f"Word export failed for {docx_path}:\n{result.stdout}\n{result.stderr}")


def measure(pdf_path):
    """Return (pages, share of the last page's usable height that text reaches)."""
    pdf = fitz.open(pdf_path)
    blocks = [b for b in pdf[-1].get_text("blocks") if b[4].strip()]
    usable_bottom = (PAGE_HEIGHT - MARGIN_BOTTOM) * 72
    return pdf.page_count, max(b[3] for b in blocks) / usable_bottom


def build(source):
    source = source.resolve()
    meta, resume = parse_source(source)
    max_pages = int(meta.get("pages", 1))
    out_docx = source.with_name(meta["output"] + ".docx")
    out_pdf = source.with_name(meta["output"] + ".pdf")

    lang = meta["lang"]
    work = Path(tempfile.mkdtemp(prefix="resume-build-"))
    tried = {}

    def attempt(size, line):
        key = f"{size}-{line}"
        if key not in tried:
            Renderer(meta, size, line).render(resume, work / f"{key}.docx")
            export_pdf(work / f"{key}.docx", work / f"{key}.pdf")
            tried[key] = measure(work / f"{key}.pdf")
        pages, fill = tried[key]
        return pages <= max_pages, fill

    try:
        index, line = SIZES.index(BASE), LINE[lang]
        fits, fill = attempt(SIZES[index], line)
        if fits:
            # Grow while the next size still fits and the page has room left.
            while fill < FILL_TARGET and index + 1 < len(SIZES):
                next_fits, next_fill = attempt(SIZES[index + 1], line)
                if not next_fits:
                    break
                index, fill = index + 1, next_fill
        else:
            while not fits and index > 0:
                index -= 1
                fits, fill = attempt(SIZES[index], line)
            # Still too long at the smallest size: tighten the line spacing down to the floor.
            floor = LEADING_FLOOR / SINGLE_LINE[lang]
            while not fits and round(line - 0.05, 2) >= floor - 0.005:
                line = round(line - 0.05, 2)
                fits, fill = attempt(SIZES[index], line)
            if not fits:
                print(f"FAILED: {source.name} does not fit {max_pages} page(s) at {SIZES[0]} pt "
                      f"with {LEADING_FLOOR}x leading; cut content")
                return False
        size = SIZES[index]
        shutil.copyfile(work / f"{size}-{line}.docx", out_docx)
        shutil.copyfile(work / f"{size}-{line}.pdf", out_pdf)
    finally:
        shutil.rmtree(work, ignore_errors=True)

    report = check(out_pdf, max_pages, meta.get("honors", "yes") != "no")
    status = "ok" if not report["problems"] else "FAILED"
    leading = line * SINGLE_LINE[lang]
    print(f"{status}: {out_pdf.name} | body {size} pt, {leading:.2f}x leading | {report['pages']} page(s) | "
          f"text reaches {fill:.0%} of the page")
    for problem in report["problems"]:
        print(f"  - {problem}")
    return not report["problems"]


def is_source(path):
    head = path.read_text(encoding="utf-8")[:400]
    return head.startswith("---") and re.search(r"^output:", head, re.M) is not None


def main():
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("sources", nargs="*", type=Path, help="Markdown resume sources")
    parser.add_argument("--all", action="store_true", help="build every source under resumes/current/")
    args = parser.parse_args()
    sources = list(args.sources)
    if args.all:
        sources += sorted(p for p in CURRENT.rglob("*.md") if is_source(p))
    if not sources:
        parser.error("give a source file or --all")
    failed = sum(not build(source) for source in sources)
    sys.exit(1 if failed else 0)


if __name__ == "__main__":
    main()
